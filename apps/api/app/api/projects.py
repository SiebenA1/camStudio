"""项目与 Brief 路由。"""
from __future__ import annotations

import io
import json
import zipfile
from datetime import datetime, timezone
from io import BytesIO
from urllib.parse import quote

from fastapi import APIRouter, Depends, File, UploadFile
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session

from app.api.deps import get_current_user
from app.core.audit import write_audit
from app.core.errors import AppError, ok
from app.core.storage import local_path
from app.database import get_db
from app.models.previz import PrevizAsset
from app.models.project import (
    Brief,
    BriefSourceType,
    Project,
    ProjectStatus,
    Variable,
    VariableCombo,
)
from app.models.reference import ReferenceShot
from app.models.review import Review
from app.models.tenant import User
from app.schemas.project import (
    BriefImportRequest,
    ComboCreate,
    ProjectCreate,
    ProjectUpdate,
    VariableUpdate,
)
from app.services.brief_service import parse_brief_to_variables

router = APIRouter(prefix="/projects", tags=["projects"])


def _project_dict(p: Project) -> dict:
    return {
        "id": p.id,
        "tenant_id": p.tenant_id,
        "name": p.name,
        "brand": p.brand,
        "status": p.status,
        "owner_id": p.owner_id,
        "budget": p.budget,
        "timeline": p.timeline,
        "description": p.description,
        "created_at": p.created_at.isoformat() if p.created_at else None,
        "updated_at": p.updated_at.isoformat() if p.updated_at else None,
    }


@router.post("")
def create_project(body: ProjectCreate, db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    p = Project(
        tenant_id=user.tenant_id,
        name=body.name,
        brand=body.brand,
        description=body.description,
        budget=body.budget,
        timeline=body.timeline,
        status=ProjectStatus.DRAFT.value,
        owner_id=user.id,
    )
    db.add(p)
    db.flush()
    write_audit(db, tenant_id=user.tenant_id, action="project.create", entity="project",
                entity_id=p.id, actor_id=user.id, actor_name=user.name, after={"name": p.name})
    db.commit()
    return ok(_project_dict(p))


@router.get("")
def list_projects(db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    projects = (
        db.query(Project)
        .filter(Project.tenant_id == user.tenant_id, Project.deleted_at.is_(None))
        .order_by(Project.updated_at.desc())
        .all()
    )
    return ok([_project_dict(p) for p in projects])


@router.get("/{project_id}")
def get_project(project_id: str, db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    p = _get_project(db, user, project_id)
    data = _project_dict(p)
    brief = db.query(Brief).filter(Brief.project_id == p.id).first()
    data["brief"] = (
        {"id": brief.id, "source_type": brief.source_type, "raw_text": brief.raw_text,
         "parsed_json": brief.parsed_json, "status": brief.status}
        if brief else None
    )
    return ok(data)


@router.patch("/{project_id}")
def update_project(project_id: str, body: ProjectUpdate, db: Session = Depends(get_db),
                   user: User = Depends(get_current_user)):
    p = _get_project(db, user, project_id)
    for k, v in body.model_dump(exclude_unset=True).items():
        setattr(p, k, v)
    db.commit()
    return ok(_project_dict(p))


@router.delete("/{project_id}")
def delete_project(project_id: str, db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    p = _get_project(db, user, project_id)
    p.deleted_at = datetime.now(timezone.utc)
    write_audit(db, tenant_id=user.tenant_id, action="project.delete", entity="project",
                entity_id=p.id, actor_id=user.id, actor_name=user.name)
    db.commit()
    return ok(None, "已删除")


@router.get("/{project_id}/export")
def export_project(project_id: str, db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    """导出项目为 ZIP：project.json（全部结构化数据）+ 空镜图 + 预演图片。"""
    p = _get_project(db, user, project_id)

    brief = db.query(Brief).filter(Brief.project_id == p.id).first()
    variables = db.query(Variable).filter(Variable.project_id == p.id).order_by(Variable.sort).all()
    combos = db.query(VariableCombo).filter(VariableCombo.project_id == p.id).all()
    shots = db.query(ReferenceShot).filter(ReferenceShot.project_id == p.id).all()
    assets = (
        db.query(PrevizAsset)
        .filter(PrevizAsset.project_id == p.id)
        .order_by(PrevizAsset.created_at)
        .all()
    )
    reviews = db.query(Review).filter(Review.project_id == p.id).all()

    data = {
        "exported_at": datetime.now(timezone.utc).isoformat(),
        "project": _project_dict(p),
        "brief": (
            {
                "source_type": brief.source_type,
                "raw_text": brief.raw_text,
                "parsed_json": brief.parsed_json,
                "status": brief.status,
            }
            if brief
            else None
        ),
        "variables": [
            {"type": v.type, "options": v.options, "selected": v.selected,
             "weight": v.weight, "sort": v.sort}
            for v in variables
        ],
        "variable_combos": [{"combo": c.combo, "weight": c.weight} for c in combos],
        "reference_shots": [
            {
                "id": s.id,
                "file": f"reference_shots/{s.id[:8]}",
                "filename": s.filename,
                "analysis_json": s.analysis_json,
                "guidance_json": s.guidance_json,
                "status": s.status,
            }
            for s in shots
        ],
        "previz_assets": [
            {
                "id": a.id,
                "file": f"assets/{a.id[:8]}",
                "model": a.model,
                "provider": a.provider,
                "prompt": a.prompt,
                "seed": a.seed,
                "params": a.params,
                "status": a.status,
                "quality_score": a.quality_score,
                "quality_detail": a.quality_detail,
                "cost": a.cost,
                "version": a.version,
                "ai_ratio": a.ai_ratio,
                "created_at": a.created_at.isoformat() if a.created_at else None,
            }
            for a in assets
        ],
        "reviews": [
            {
                "id": r.id,
                "title": r.title,
                "type": r.type,
                "status": r.status,
                "decision": r.decision,
                "asset_ids": r.asset_ids,
                "comments": [
                    {"author": c.author_name, "text": c.text, "score": c.score,
                     "created_at": c.created_at.isoformat() if c.created_at else None}
                    for c in r.comments
                ],
                "votes": [{"asset_id": v.asset_id, "vote": v.vote} for v in r.votes],
                "approvals": [
                    {"actor": ap.actor_name, "decision": ap.decision, "comment": ap.comment}
                    for ap in r.approvals
                ],
            }
            for r in reviews
        ],
    }

    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("project.json", json.dumps(data, ensure_ascii=False, indent=2))
        for a in assets:
            lp = local_path(a.file_url)
            if lp and lp.exists():
                zf.write(lp, f"assets/{a.id[:8]}{lp.suffix}")
        for s in shots:
            lp = local_path(s.file_url)
            if lp and lp.exists():
                zf.write(lp, f"reference_shots/{s.id[:8]}{lp.suffix}")
    buf.seek(0)

    filename = f"realframe-{p.name}.zip"
    return StreamingResponse(
        buf,
        media_type="application/zip",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(filename)}"},
    )


# ---- Brief ----
@router.post("/{project_id}/brief")
def import_brief(project_id: str, body: BriefImportRequest, db: Session = Depends(get_db),
                 user: User = Depends(get_current_user)):
    _get_project(db, user, project_id)
    brief = db.query(Brief).filter(Brief.project_id == project_id).first()
    if not brief:
        brief = Brief(project_id=project_id, source_type=body.source_type, raw_text=body.raw_text)
        db.add(brief)
    else:
        brief.source_type = body.source_type
        brief.raw_text = body.raw_text
        brief.status = "pending"
    db.commit()
    return ok({"id": brief.id, "raw_text": brief.raw_text, "source_type": brief.source_type})


@router.post("/{project_id}/brief/file")
async def import_brief_file(project_id: str, file: UploadFile = File(...),
                            db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    _get_project(db, user, project_id)
    content = await file.read()
    text = _extract_file_text(file.filename or "", content)
    source_type = BriefSourceType.PDF.value if (file.filename or "").lower().endswith(".pdf") else (
        BriefSourceType.PPT.value if (file.filename or "").lower().endswith((".ppt", ".pptx")) else BriefSourceType.TEXT.value
    )
    brief = db.query(Brief).filter(Brief.project_id == project_id).first()
    if not brief:
        brief = Brief(project_id=project_id, source_type=source_type, raw_text=text)
        db.add(brief)
    else:
        brief.source_type = source_type
        brief.raw_text = text
        brief.status = "pending"
    db.commit()
    return ok({"id": brief.id, "raw_text": text, "source_type": source_type})


@router.post("/{project_id}/brief/parse")
async def parse_brief(project_id: str, db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    _get_project(db, user, project_id)
    brief = db.query(Brief).filter(Brief.project_id == project_id).first()
    if not brief or not brief.raw_text.strip():
        raise AppError(400, "no_brief", "请先导入 Brief 内容")
    data = await parse_brief_to_variables(db, user.tenant_id, brief)
    return ok(data)


# ---- 变量 ----
@router.get("/{project_id}/variables")
def list_variables(project_id: str, db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    _get_project(db, user, project_id)
    vs = db.query(Variable).filter(Variable.project_id == project_id).order_by(Variable.sort).all()
    return ok([
        {"id": v.id, "type": v.type, "options": v.options, "selected": v.selected,
         "weight": v.weight, "sort": v.sort} for v in vs
    ])


@router.patch("/{project_id}/variables/{vid}")
def update_variable(project_id: str, vid: str, body: VariableUpdate,
                    db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    _get_project(db, user, project_id)
    v = db.get(Variable, vid)
    if not v or v.project_id != project_id:
        raise AppError(404, "not_found", "变量不存在")
    for k, val in body.model_dump(exclude_unset=True).items():
        setattr(v, k, val)
    db.commit()
    return ok({"id": v.id, "type": v.type, "options": v.options, "selected": v.selected,
               "weight": v.weight, "sort": v.sort})


@router.post("/{project_id}/combos")
def create_combo(project_id: str, body: ComboCreate, db: Session = Depends(get_db),
                 user: User = Depends(get_current_user)):
    _get_project(db, user, project_id)
    c = VariableCombo(project_id=project_id, combo=body.combo, weight=body.weight)
    db.add(c)
    db.commit()
    return ok({"id": c.id, "combo": c.combo, "weight": c.weight})


# ---- helpers ----
def _get_project(db: Session, user: User, project_id: str) -> Project:
    p = db.get(Project, project_id)
    if not p or p.deleted_at is not None:
        raise AppError(404, "not_found", "项目不存在")
    if p.tenant_id != user.tenant_id:
        raise AppError(403, "forbidden", "无权访问该项目")
    return p


def _extract_file_text(filename: str, content: bytes) -> str:
    lower = filename.lower()
    try:
        if lower.endswith(".pdf"):
            from pypdf import PdfReader
            reader = PdfReader(BytesIO(content))
            return "\n".join(page.extract_text() or "" for page in reader.pages).strip()
        if lower.endswith((".ppt", ".pptx")):
            from pptx import Presentation
            prs = Presentation(BytesIO(content))
            parts = []
            for slide in prs.slides:
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text:
                        parts.append(shape.text)
            return "\n".join(parts).strip()
    except Exception as e:
        raise AppError(400, "extract_failed", f"文件解析失败：{e}")
    return content.decode("utf-8", errors="ignore").strip()
